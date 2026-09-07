#!/usr/bin/env python3
"""
Importe les niveaux de risque des FCP depuis le classeur « Classement des
fonds ».

Entree  : Classement des fonds au JJMMAAAA.xlsx
Sortie  : data/fcp/risque.csv
          gestionnaire;nomAumfcp;niveauRisque;echelleMax;horizonAnnees;source;dateReleve

Le classeur melange DEUX dispositions, et les confondre fait attribuer des
niveaux a des fonds qui n'en ont pas :

  - Feuilles « Actions » et « Diversifies » : les fonds sont groupes sous des
    lignes de titre « Niveau de risque N ». Un de ces groupes s'intitule
    « Niveau de risque non determine » — un motif qui ne cherche qu'un chiffre
    ne le reconnait pas, et le niveau du groupe precedent se reporte
    silencieusement sur sept fonds.

  - Feuilles « Obligations », « Monetaire » et « Halal » : le niveau est une
    COLONNE, pas un groupe, avec « N/A » quand il est inconnu.

Le rapprochement avec le referentiel se fait sur le NOM DU FONDS et non sur le
gestionnaire : le classeur emploie ses propres libelles de societes
(« BOA ASSET MANAGEMENT » contre « BOA CAPITAL CAM », « SOGESPAR » contre
« SGCAMWA », « OAM S.A. » contre « OPTI ASSET MANAGEMENT »). Le gestionnaire ne
sert qu'a signaler une incoherence.

Usage :
    python scripts/import_fcp_risque.py "Classement des fonds au 30012026 (1).xlsx"
    python scripts/import_fcp_risque.py <classeur> --dry-run
"""

from __future__ import annotations

import argparse
import csv
import os
import re
import sys
import unicodedata

try:
    import openpyxl
except ImportError:
    print("openpyxl requis : python -m pip install openpyxl", file=sys.stderr)
    raise SystemExit(1)

RACINE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
AUMFCP = os.path.join(RACINE, "data", "fcp", "aumfcp.csv")
SORTIE = os.path.join(RACINE, "data", "fcp", "risque.csv")

ENTETES = [
    "gestionnaire",
    "nomAumfcp",
    "niveauRisque",
    "echelleMax",
    "horizonAnnees",
    "source",
    "dateReleve",
]

# Le classeur note les niveaux sur l'echelle reglementaire des DICI.
ECHELLE_MAX = 7

# Libelles du classeur a ramener au nom du referentiel. Chacun verifie a la
# main : meme gestionnaire, meme categorie, meme VL d'origine. Ce sont des
# ecarts de nommage, pas des fonds distincts.
ALIAS = {
    # Le classeur numerote la forme juridique, le referentiel numerote le fonds.
    "FCP-1 OPTI PLACEMENT": "OPTI PLACEMENT 1",
    "FCP-2 OPTI REVENU": "OPTI REVENU 2",
    "FCP-3 OPTI CAPITAL": "OPTI CAPITAL 3",
    # Le classeur porte l'acronyme de la societe, le referentiel non.
    "FCP AAM EPARGNE ACTION": "EPARGNE ACTION",
    "FCP AAM EPARGNE CROISSANCE": "EPARGNE CROISSANCE",
    # L'inverse ici : le referentiel porte l'acronyme, pas le classeur.
    "FCP CAPITAL PLUS": "AGA CAPITAL PLUS",
    "FCP CONFORT PLUS": "AGA CONFORT PLUS",
    # Abreviations du referentiel.
    "FCPE ORANGE MALI": "ORANGE ML",
    "FCP IMPACT DIASPORA": "DIASPORA",
}

# Valeurs signifiant « inconnu ». On ne les importe pas : une case vide dit la
# meme chose sans faire croire a une donnee.
INCONNU = {"", "N/A", "NA", "ND", "-", "NON DETERMINE", "NON DÉTERMINÉ"}


def sans_accent(s: str) -> str:
    return unicodedata.normalize("NFD", s).encode("ascii", "ignore").decode()


def cle(raw) -> str:
    """Cle de rapprochement d'un nom de fonds : sans forme juridique ni accent."""
    k = sans_accent(str(raw or "").strip()).upper()
    k = re.sub(r"^(FCP|FCPR|FCTC|FCC|FCPE|SICAV|SICAF)[-\s]+", "", k)
    # « FCP-1 OPTI PLACEMENT », « FCP-2 OPTI REVENU » : le classeur numerote la
    # forme juridique, pas le fonds.
    k = re.sub(r"^\d+[\s-]+", "", k)
    return re.sub(r"[^A-Z0-9]", "", k)


def charger_referentiel() -> dict[str, tuple[str, str]]:
    """{cle du nom: (gestionnaire, nom canonique)} depuis aumfcp.csv."""
    out: dict[str, tuple[str, str]] = {}
    with open(AUMFCP, encoding="cp1252", newline="") as f:
        for r in csv.DictReader(f, delimiter=";"):
            gest = (r.get("Gestionnaire") or "").strip()
            nom = (r.get("Nom de l'OPC") or "").strip()
            if gest and nom:
                out[cle(nom)] = (gest, nom)
    return out


def extraire(chemin: str) -> list[dict]:
    """Lit le classeur et rend une ligne par fonds, niveau compris ou None."""
    wb = openpyxl.load_workbook(chemin, data_only=True, read_only=True)
    lignes: list[dict] = []

    for feuille in wb.sheetnames:
        ws = wb[feuille]
        niveau_groupe: int | None = None
        colonne_niveau: int | None = None

        for brut in ws.iter_rows(min_row=1, values_only=True):
            cells = [
                "" if c is None else str(c).strip() for c in brut
            ]
            if not any(cells):
                continue

            # En-tete de tableau : on y cherche une colonne « NIVEAU DE RISQUE ».
            if any(c.upper() == "NOM SGO" for c in cells):
                colonne_niveau = None
                for i, c in enumerate(cells):
                    if "NIVEAU DE RISQUE" in sans_accent(c).upper():
                        colonne_niveau = i
                continue

            # Ligne de groupe « Niveau de risque N » ou « ... non determine ».
            joint = sans_accent(" | ".join(cells)).upper()
            if "NIVEAU DE RISQUE" in joint:
                m = re.search(r"NIVEAU DE RISQUE\s*:?\s*(\d+)", joint)
                # Sans chiffre — « non determine » — le groupe REMET a None.
                # C'est tout l'objet de cette branche : sans elle, le niveau
                # precedent se reporterait.
                niveau_groupe = int(m.group(1)) if m else None
                continue

            sgo = cells[2] if len(cells) > 2 else ""
            produit = cells[3] if len(cells) > 3 else ""
            if not sgo or not produit:
                continue

            niveau = niveau_groupe
            if colonne_niveau is not None and len(cells) > colonne_niveau:
                v = cells[colonne_niveau].strip()
                if sans_accent(v).upper() in INCONNU:
                    niveau = None
                else:
                    try:
                        niveau = int(float(v.replace(",", ".")))
                    except ValueError:
                        niveau = None

            lignes.append(
                {
                    "feuille": feuille,
                    "sgo": sgo,
                    "produit": produit,
                    "niveau": niveau,
                }
            )
    return lignes


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("classeur")
    ap.add_argument("--dry-run", action="store_true")
    args = ap.parse_args()

    if not os.path.exists(args.classeur):
        print(f"Classeur introuvable : {args.classeur}", file=sys.stderr)
        return 1

    ref = charger_referentiel()
    lignes = extraire(args.classeur)
    date_releve = re.search(r"(\d{2})(\d{2})(\d{4})", os.path.basename(args.classeur))
    date_iso = (
        f"{date_releve.group(3)}-{date_releve.group(2)}-{date_releve.group(1)}"
        if date_releve
        else ""
    )

    print(f"Classeur : {os.path.basename(args.classeur)}")
    print(f"{len(lignes)} fonds lus, relevé au {date_iso or 'date inconnue'}")

    avec = [l for l in lignes if l["niveau"] is not None]
    sans = [l for l in lignes if l["niveau"] is None]
    print(f"  niveau renseigné : {len(avec)}")
    print(f"  niveau inconnu   : {len(sans)}")

    sorties: list[list[str]] = []
    orphelins: list[dict] = []
    desaccords: list[tuple[str, str, str]] = []

    for l in avec:
        nom_vise = ALIAS.get(l["produit"].strip().upper(), l["produit"])
        cible = ref.get(cle(nom_vise))
        if cible is None:
            orphelins.append(l)
            continue
        gest, nom = cible
        if cle(l["sgo"])[:8] != cle(gest)[:8]:
            desaccords.append((nom, l["sgo"], gest))
        sorties.append(
            [
                gest,
                nom,
                str(l["niveau"]),
                str(ECHELLE_MAX),
                "",  # horizon absent du classeur
                f"Classement des fonds au {date_iso} (feuille {l['feuille']})",
                date_iso,
            ]
        )

    print(f"\n{len(sorties)} rapprochés au référentiel")
    if orphelins:
        print(f"\n{len(orphelins)} fonds du classeur absents du référentiel :")
        for l in orphelins:
            print(f"   niveau {l['niveau']}  {l['produit'][:38]:40} {l['sgo'][:26]}")
    if desaccords:
        print(f"\n{len(desaccords)} libellés de gestionnaire différents (sans gravité,")
        print("  le rapprochement se fait sur le nom du fonds) :")
        for nom, a, b in desaccords:
            print(f"   {nom[:26]:28} classeur « {a[:24]} » ≠ référentiel « {b[:24]} »")
    if sans:
        print(f"\n{len(sans)} fonds sans niveau dans le classeur — à chercher ailleurs :")
        for l in sans:
            print(f"   {l['feuille']:12} {l['produit'][:38]:40} {l['sgo'][:26]}")

    if args.dry_run:
        print("\n--dry-run : rien n'a été écrit.")
        return 0

    # Les entrees deja presentes (relevees a la main ou sur les sites des SGO)
    # sont conservees si le classeur ne couvre pas le fonds.
    existant: dict[str, list[str]] = {}
    if os.path.exists(SORTIE):
        with open(SORTIE, encoding="utf-8", newline="") as f:
            for r in csv.DictReader(f, delimiter=";"):
                existant[cle(r["nomAumfcp"])] = [r[c] for c in ENTETES]

    for s in sorties:
        existant[cle(s[1])] = s

    with open(SORTIE, "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(ENTETES)
        w.writerows(sorted(existant.values(), key=lambda r: (r[0], r[1])))

    print(f"\n{len(existant)} fonds dans data/fcp/risque.csv")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
